import io
import pandas as pd
from datetime import timedelta
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from flask import Flask, request, send_file, render_template_string

# ---------------------------------------------------------------
# 0. ВЕБ-ПРИЛОЖЕНИЕ (вместо пути к папке - загрузка файлов с сайта)
# ---------------------------------------------------------------
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024   # не больше 20 МБ за раз

RESULT_NAME = "результат_проверки.xlsx"   # имя файла, который скачает пользователь

# ---------------------------------------------------------------
# 2. НАСТРОЙКИ ПОРОГОВ
# ---------------------------------------------------------------
CASH_LIMIT = 1_000_000        # операции с наличными от этой суммы - обязательный контроль
BIG_AMOUNT = 300_000          # крупная операция
SMALL_AMOUNT = 100_000        # мелкая операция

# Правило 6: (период в часах, минимум операций, минимальная общая сумма)
MANY_OPS_RULES = [
    (1, 3, 300_000),          # 3 операции за час на 300 000 и больше
    (24 * 7, 5, 300_000),     # 5 операций за неделю на 300 000 и больше
]

MARK_MANDATORY = "ОБЯЗАТЕЛЬНЫЙ КОНТРОЛЬ"
MARK_SUSPICIOUS = "ПОДОЗРИТЕЛЬНАЯ"
MARK_BOTH = "ОБЯЗАТЕЛЬНЫЙ КОНТРОЛЬ + ПОДОЗРИТЕЛЬНАЯ"
MARK_OK = "ВСЁ НОРМАЛЬНО"
PRIORITY = {MARK_OK: 0, MARK_SUSPICIOUS: 1, MARK_MANDATORY: 2, MARK_BOTH: 3}

REQUIRED_COLUMNS = ["ФИО", "Дата и время", "Сумма", "Тип операции"]

# Необязательные столбцы: если их нет в файле, соответствующее правило не срабатывает
CASH_COLUMN = "Способ"        # "Наличные" или "Безналичные"
COUNTRY_COLUMN = "Страна"     # страна, откуда пришёл перевод или куда он ушёл

CASH_VALUES = {"наличные", "наличными", "нал", "cash"}

# Страны из "чёрного списка" ФАТФ (не выполняют меры против отмывания
# денег и финансирования терроризма). Список можно дополнять.
HIGH_RISK_COUNTRIES = {
    "иран", "iran",
    "кндр", "северная корея", "корейская народно-демократическая республика",
    "north korea", "dprk",
    "мьянма", "бирма", "myanmar",
}


def analyze(tables):
    """Вся проверка из исходного скрипта. На входе - список таблиц,
    на выходе - готовый Excel-файл в памяти."""

    # -----------------------------------------------------------
    # 3. ОБЪЕДИНЕНИЕ ЗАГРУЖЕННЫХ ТАБЛИЦ
    # -----------------------------------------------------------
    df = pd.concat(tables, ignore_index=True)
    df["Дата и время"] = pd.to_datetime(df["Дата и время"], dayfirst=True)
    df = df.sort_values(["ФИО", "Дата и время"]).reset_index(drop=True)

    # Если необязательных столбцов нет - создаём пустые, чтобы код не падал
    for column in (CASH_COLUMN, COUNTRY_COLUMN):
        if column not in df.columns:
            df[column] = ""
    df[CASH_COLUMN] = df[CASH_COLUMN].fillna("")
    df[COUNTRY_COLUMN] = df[COUNTRY_COLUMN].fillna("")

    df["Маркер"] = MARK_OK
    df["Причина"] = ""

    def set_marker(index, marker, reason):
        """Ставит маркер, только если он строже текущего.
        Обязательный контроль + подозрительная = общий маркер MARK_BOTH."""
        current = df.at[index, "Маркер"]
        if {current, marker} == {MARK_MANDATORY, MARK_SUSPICIOUS}:
            df.at[index, "Маркер"] = MARK_BOTH
        elif PRIORITY[marker] > PRIORITY[current]:
            df.at[index, "Маркер"] = marker
        if reason not in df.at[index, "Причина"]:
            df.at[index, "Причина"] += reason + "; "

    def check_series(big_ops, hours, count, text):
        """Несколько крупных операций за короткое время."""
        for i, op in big_ops.iterrows():
            end = op["Дата и время"] + timedelta(hours=hours)
            group = big_ops[(big_ops["Дата и время"] >= op["Дата и время"]) &
                            (big_ops["Дата и время"] <= end)]
            if len(group) >= count:
                for j in group.index:
                    set_marker(j, MARK_SUSPICIOUS, f"{text} ({len(group)} шт.)")

    # -----------------------------------------------------------
    # 4. АНАЛИЗ КАЖДОГО ЧЕЛОВЕКА
    # -----------------------------------------------------------
    for name, person in df.groupby("ФИО"):

        # Правило 1: операция с наличными от 1 000 000
        for i, op in person.iterrows():
            is_cash = str(op[CASH_COLUMN]).strip().lower() in CASH_VALUES
            if is_cash and op["Сумма"] >= CASH_LIMIT:
                set_marker(i, MARK_MANDATORY, "Операция с наличными от 1 000 000 руб.")

        # Правило 1.1: перевод из страны или в страну из чёрного списка ФАТФ
        for i, op in person.iterrows():
            country = str(op[COUNTRY_COLUMN]).strip().lower()
            if country in HIGH_RISK_COUNTRIES:
                set_marker(i, MARK_MANDATORY,
                           f"Перевод из/в страну высокого риска ({op[COUNTRY_COLUMN]})")

        big = person[person["Сумма"] >= BIG_AMOUNT]

        # Правило 2: несколько крупных за 1 час
        check_series(big, hours=1, count=2,
                     text="Несколько крупных операций за час")

        # Правило 3: несколько крупных за сутки
        check_series(big, hours=24, count=3,
                     text="Несколько крупных операций за сутки")

        # Правило 4: крупное поступление раздроблено на мелкие списания
        for i, op in big.iterrows():
            if op["Тип операции"] != "Поступление":
                continue
            end = op["Дата и время"] + timedelta(hours=24)
            small_out = person[(person["Тип операции"] == "Списание") &
                               (person["Дата и время"] > op["Дата и время"]) &
                               (person["Дата и время"] <= end) &
                               (person["Сумма"] < SMALL_AMOUNT)]
            if len(small_out) >= 3 and small_out["Сумма"].sum() >= 0.7 * op["Сумма"]:
                text = "Крупное поступление раздроблено на мелкие списания"
                set_marker(i, MARK_SUSPICIOUS, text)
                for j in small_out.index:
                    set_marker(j, MARK_SUSPICIOUS, text)

        # Правило 5: крупная операция, а до неё по счёту не было никаких операций
        for i, op in big.iterrows():
            earlier = person[person["Дата и время"] < op["Дата и время"]]
            if len(earlier) == 0:
                set_marker(i, MARK_SUSPICIOUS,
                           "Крупная операция, ранее операций по счёту не было")

        # Правило 6: много операций (даже мелких) за период на крупную сумму
        for hours, min_count, min_sum in MANY_OPS_RULES:
            for i, op in person.iterrows():
                end = op["Дата и время"] + timedelta(hours=hours)
                group = person[(person["Дата и время"] >= op["Дата и время"]) &
                               (person["Дата и время"] <= end)]
                if len(group) >= min_count and group["Сумма"].sum() >= min_sum:
                    for j in group.index:
                        set_marker(j, MARK_SUSPICIOUS,
                                   f"Много операций за {hours} ч на крупную сумму")

    # -----------------------------------------------------------
    # 5. ИТОГ ПО КЛИЕНТАМ (будет на втором листе Excel)
    # -----------------------------------------------------------
    summary_rows = []
    for name, person in df.groupby("ФИО"):
        both = (person["Маркер"] == MARK_BOTH).sum()
        mandatory = (person["Маркер"] == MARK_MANDATORY).sum()
        suspicious = (person["Маркер"] == MARK_SUSPICIOUS).sum()
        if both > 0:
            verdict = ("есть операции обязательного контроля "
                       "с признаками подозрительности")
        elif mandatory > 0:
            verdict = "есть операции обязательного контроля"
        elif suspicious > 0:
            verdict = "есть подозрительные операции"
        else:
            verdict = "странных операций нет"
        summary_rows.append({"ФИО": name,
                             "Обяз. контроль + подозрительная": both,
                             "Обязательный контроль": mandatory,
                             "Подозрительных": suspicious,
                             "Вердикт": verdict})
    summary = pd.DataFrame(summary_rows)

    # -----------------------------------------------------------
    # 6. СОХРАНЕНИЕ В EXCEL С ЦВЕТАМИ (в память, а не на диск)
    # -----------------------------------------------------------
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Операции", index=False)
        summary.to_excel(writer, sheet_name="Итог по клиентам", index=False)

    buffer.seek(0)
    wb = load_workbook(buffer)
    ws = wb["Операции"]
    colors = {
        MARK_BOTH: PatternFill("solid", fgColor="D9B8F0"),        # фиолетовый
        MARK_MANDATORY: PatternFill("solid", fgColor="FFC7CE"),   # красный
        MARK_SUSPICIOUS: PatternFill("solid", fgColor="FFEB9C"),  # жёлтый
        MARK_OK: PatternFill("solid", fgColor="C6EFCE"),          # зелёный
    }
    marker_col = list(df.columns).index("Маркер") + 1
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=marker_col)
        cell.fill = colors[cell.value]

    result = io.BytesIO()
    wb.save(result)
    result.seek(0)
    return result


# ---------------------------------------------------------------
# 7. СТРАНИЦА САЙТА
# ---------------------------------------------------------------
PAGE = """
<!doctype html>
<html lang="ru">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Мониторинг операций</title>
  <style>
    body { font-family: Arial, sans-serif; max-width: 560px; margin: 60px auto; padding: 0 16px; }
    button { padding: 10px 20px; font-size: 16px; cursor: pointer; }
    .error { color: #b00020; }
  </style>
</head>
<body>
  <h1>Мониторинг операций</h1>
  <p>Загрузите один или несколько Excel-файлов (.xlsx) со столбцами:
     ФИО, Дата и время, Сумма, Тип операции.
     Дополнительно можно добавить столбцы Способ (Наличные / Безналичные) и Страна.</p>
  <form method="post" action="/check" enctype="multipart/form-data">
    <p><input type="file" name="files" accept=".xlsx,.xlsm" multiple required></p>
    <p><button type="submit">Проверить и скачать результат</button></p>
  </form>
  {% if error %}<p class="error">{{ error }}</p>{% endif %}
</body>
</html>
"""


@app.route("/")
def index():
    return render_template_string(PAGE)


@app.route("/check", methods=["POST"])
def check():
    # Загрузка файлов, которые прислал пользователь (вместо чтения папки)
    tables = []
    for file in request.files.getlist("files"):
        if not file.filename.lower().endswith((".xlsx", ".xlsm")):
            continue
        try:
            table = pd.read_excel(file)
        except Exception:
            return render_template_string(PAGE, error=f"Не удалось прочитать файл {file.filename}")
        missing = [c for c in REQUIRED_COLUMNS if c not in table.columns]
        if missing:
            return render_template_string(
                PAGE, error=f"В файле {file.filename} нет столбцов: {', '.join(missing)}")
        table["Файл"] = file.filename
        tables.append(table)

    if not tables:
        return render_template_string(PAGE, error="Не загружено ни одного Excel-файла (.xlsx)")

    try:
        result = analyze(tables)
    except Exception as e:
        return render_template_string(PAGE, error=f"Ошибка при обработке: {e}")

    return send_file(result, as_attachment=True, download_name=RESULT_NAME,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app.run(debug=True)
